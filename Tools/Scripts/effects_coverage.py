"""
Effects coverage workbook: which new cards' effects WolfX can read (docs/Card Lists/Effects Coverage.xlsx).

usage: python effects_coverage.py <translated cards.json> <work.csv> <out.xlsx>

Both inputs come from one `WolfX.exe --attach-effects <cards.json> <report.txt> <work.csv> <mod folder>` run: cards.json has the
cards whose whole text was translated (effectScript / effectClone attached, effectLooser = what is approximated), work.csv lists
every card with an effect still to write (effects read, the sentences not run and why). Three sheets, each an Excel table:
Fully read, Partially read, Not read - with a "why" column in plain words.
"""
import csv
import json
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.table import Table, TableStyleInfo

# The translator's reason prefixes ("why not run" in work.csv) in plain words.
REASONS = [
    ("head", "Condition or timing not read yet"),
    ("action", "Action not supported yet"),
    ("cost", "Cost not supported yet"),
    ("rule", "Continuous rule not supported yet"),
    ("selector", "Card filter not supported yet"),
    ("bullet", "Bullet under a choice / condition that cannot run yet"),
    ("negate", "Negation form not supported yet"),
    ("follow-up", "Only removes your own card for a follow-up that is not read"),
    ("spell/trap with a trigger", "Spell/Trap effect with a trigger"),
    ("self summon with a trigger", "Special Summon this card on a trigger"),
    ("self summon with cost", "Special Summon this card with a self cost"),
    ("chain response that is not a negation", "Chain response other than a negation"),
    ("equip spell effect other than equip", "Equip Spell effect other than the equip"),
    ("battle phase effect that is not a Quick Effect", "Battle Phase effect that is not a Quick Effect"),
    ("conflict", "Two parts need the effect's one filter row"),
]


def explain(reason: str) -> str:
    reason = reason.strip()
    for prefix, words in REASONS:
        if reason.startswith(prefix + ":") or reason == prefix:
            rest = reason[len(prefix) + 1:].strip()
            return f"{words}: {rest}" if rest else words
    return reason


def table(ws, name, header, rows, widths):
    ws.append(header)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + i)].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    last = chr(ord("A") + len(header) - 1)
    t = Table(displayName=name, ref=f"A1:{last}{max(2, len(rows) + 1)}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)


def main():
    cards_path, work_path, out_path = sys.argv[1], sys.argv[2], sys.argv[3]
    cards = json.load(open(cards_path, encoding="utf-8"))
    cards = cards["cards"] if isinstance(cards, dict) else cards
    work = list(csv.DictReader(open(work_path, encoding="utf-8")))
    listed = {r["id"] for r in work}

    full = []
    for c in cards:
        if str(c.get("id")) in listed or not c.get("effectScript"):
            continue
        looser = c.get("effectLooser") or ""
        why = "Every sentence is translated and the effect is attached."
        if looser:
            why += " Approximated: " + looser
        effects = len([e for e in re.split(r" also | or ", c["effectScript"]) if e.strip()])
        full.append([int(c["id"]), c.get("name", ""), c.get("kind", ""), effects, why, c["effectScript"]])

    partial, none = [], []
    for r in work:
        reasons = [w for w in r["why not run"].split(" | ") if w.strip()]
        why = "\n".join("- " + explain(w) for w in reasons) or "No reason recorded"
        row = [int(r["id"]), r["name"], r["kind"], int(r["effects read"] or 0), len(reasons), why, r["not run"].replace(" | ", "\n"), r["script so far"]]
        (partial if row[3] > 0 else none).append(row)

    # most nearly finished first
    full.sort(key=lambda x: x[0])
    partial.sort(key=lambda x: (x[4], x[0]))
    none.sort(key=lambda x: (x[4], x[0]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Fully read"
    table(ws, "FullyRead", ["id", "name", "kind", "effects", "why", "script"], full, [8, 34, 12, 8, 70, 80])
    table(wb.create_sheet("Partially read"), "PartiallyRead",
          ["id", "name", "kind", "effects read", "sentences left", "why", "sentences not read", "script so far"], partial, [8, 34, 12, 8, 9, 70, 70, 60])
    table(wb.create_sheet("Not read"), "NotRead",
          ["id", "name", "kind", "effects read", "sentences left", "why", "sentences not read", "script so far"], none, [8, 34, 12, 8, 9, 70, 70, 20])
    wb.save(out_path)
    print(f"fully read {len(full)}, partially read {len(partial)}, not read {len(none)} -> {out_path}")


if __name__ == "__main__":
    main()
