import json, collections, sys, re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
sys.stdout.reconfigure(encoding='utf-8')
raw = json.load(open(r'C:\Users\Johnathon\Desktop\New folder\effect_tables_raw.json'))
game = json.load(open(r'C:\Users\Johnathon\Desktop\New folder\Cards\Game Cards.json', encoding='utf-8-sig')); game = game['cards'] if isinstance(game, dict) else game
info = {c['id']: c for c in game}
idt = json.load(open(r'C:\Users\Johnathon\Desktop\New folder\id_tables.json'))
known = {  # functions named in IDA so far
 '0x14015f9c0': 'Effect_DrawCards_Resolve', '0x1400fac40': 'Effect_DrawCards_CanActivate', '0x140156e70': 'Effect_RunSlot1Action_Driver',
 '0x140180f10': 'Target_Raigeki_OpponentMonsters', '0x1401952d0': 'Target_DarkHole_AllMonsters', '0x1400de060': 'Slot_ReturnConst2',
 '0x1400fe620': 'Slot_Delegate_ToSecondaryTable_FE560', '0x1401f94c0': 'Slot0_DefaultMonsterResolve_1F94C0'}
fmt = lambda p: (known.get(p) and f'{known[p]} ({p})') or (p or '-')
wb = Workbook(); ws = wb.active; ws.title = 'Read me'
for line in [
 'Effect tables of YuGiOh.exe (decoded 2026-09-28). Every effect lookup goes through YGO::Effects::Get_EffectTableEntryForCard (0x1400DFBC0).',
 'Four main tables of 48-byte rows {u16 id, 3 x u16 extra, 5 function pointers (slot0..slot4)}, sorted by id; several consecutive rows = several effects.',
 'slot0 = resolve/execute, slot1 = target predicate (player, zone), slot2 = condition check (empty = allowed); slot3/slot4 not decoded yet.',
 'An EFFECT TYPE = rows with an identical set of five function pointers: the same implementation. Clone a vanilla card of the right type to give a custom card that behaviour.',
 'Sheets: Tables (all 105 id-keyed tables found), then per main table the effect types by size with example cards + their text, then By card.',
 'Names of functions come from IDA (YGO::Effects::...) as they get identified; unnamed ones show only the address.']:
    ws.append([line])
ws.column_dimensions['A'].width = 150
ws2 = wb.create_sheet('Tables'); ws2.append(['address', 'row size', 'rows', 'kind'])
main = {0x140B38290: 'main table kind 3 (head)', 0x140B39B20: 'main table kind 3 (body)', 0x140B5CC20: 'main table kind 1 (monster effects)', 0x140B85A50: 'main table kind 2', 0x140BA87C0: 'main table kind 0 (spell/trap, from its first id-3000 defaults @140BA8730)'}
for t in idt:
    a = int(t['addr'], 16); ws2.append([t['addr'], t['stride'], t['rows'], main.get(a, 'has function pointers: secondary/per-card handler table' if t['stride'] >= 16 else 'pure data (card id -> value / member list)')])
short = lambda s, n=170: re.sub(r'\s+', ' ', s or '')[:n]
byname = {}
for tn, rows in raw.items():
    clusters = collections.defaultdict(list)
    for r in rows: clusters[tuple(r['fn'])].append(r)
    ws3 = wb.create_sheet(re.sub(r'[^A-Za-z0-9 ()]', '', tn.split('@')[0]).strip()[:31]); ws3.append(['type #', 'cards (rows)', 'slot0 resolve', 'slot1 target', 'slot2 condition', 'slot3', 'slot4', 'example cards', 'example text'])
    for c in ws3[1]: c.font = Font(bold=True)
    for n, (tup, rs) in enumerate(sorted(clusters.items(), key=lambda kv: -len(kv[1])), 1):
        ex = [info[r['id']] for r in rs if r['id'] in info][:4]
        ws3.append([n, len(rs)] + [fmt(p) for p in tup] + ['; '.join(f"{c['name']} ({c['id']})" for c in ex), short(ex[0]['description']) if ex else ''])
        for r in rs: byname.setdefault(r['id'], {})[tn.split()[0]] = n
    for i, w in enumerate([7, 10, 34, 34, 34, 18, 18, 60, 90], 1): ws3.column_dimensions[get_column_letter(i)].width = w
    ws3.freeze_panes = 'A2'
ws4 = wb.create_sheet('By card'); ws4.append(['id', 'name', 'kind', 'T5 spell/trap type #', 'T2 monster type #', 'T4 type #', 'T1 type #', 'text'])
for cid in sorted(byname):
    c = info.get(cid)
    if not c: continue
    b = byname[cid]; ws4.append([cid, c['name'], c['kind'], b.get('T5'), b.get('T2'), b.get('T4'), b.get('T1'), short(c['description'], 200)])
for i, w in enumerate([7, 34, 18, 10, 10, 10, 10, 120], 1): ws4.column_dimensions[get_column_letter(i)].width = w
ws4.freeze_panes = 'A2'
wb.save(r'C:\Users\Johnathon\Desktop\New folder\effect_tables.xlsx')
print('saved; cards with any effect row:', len(byname))
